"""Workbook ingest: Excel/CSV -> Parquet -> DuckDB, cached by file digest.

Ingest is streaming and memory bounded, so a 1 lakh (100k) row workbook - or a
much larger one - is converted without ever holding the whole sheet in memory:

1. rows are pulled from ``openpyxl`` in read-only mode, chunk by chunk;
2. each chunk is written straight to a temporary all-text Parquet file;
3. DuckDB then decides a real type per column (``TRY_CAST`` statistics over the
   whole column, not a sample) and rewrites the file as typed, ZSTD Parquet.

The typed Parquet is what every later stage reads. Nothing else in the pipeline
ever re-opens the spreadsheet.
"""
from __future__ import annotations

import logging
import os
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator, Sequence

import duckdb
import pyarrow as pa
import pyarrow.parquet as pq

from .cache import WorkbookCache, file_digest, safe_name

logger = logging.getLogger(__name__)

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
SUPPORTED_SUFFIXES = EXCEL_SUFFIXES | CSV_SUFFIXES
ROW_ID = "row_id"
DEFAULT_CHUNK_ROWS = 20_000
DEFAULT_MAX_ROWS = 5_000_000
_RESERVED = {ROW_ID}


class IngestError(ValueError):
    """Raised when a workbook cannot be converted into queryable tables."""


def quote_identifier(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def quote_literal(value: str) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def sanitize_identifier(name: str, taken: set[str], fallback: str = "column") -> str:
    """Turn an arbitrary sheet/column label into a unique SQL identifier."""
    cleaned = "".join(character if character.isalnum() else "_" for character in str(name).strip().lower())
    cleaned = "_".join(part for part in cleaned.split("_") if part)
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"{fallback}_{cleaned}" if cleaned else fallback
    if cleaned in _RESERVED:
        cleaned = f"{cleaned}_col"
    candidate, suffix = cleaned[:120], 2
    while candidate in taken:
        candidate = f"{cleaned[:115]}_{suffix}"
        suffix += 1
    taken.add(candidate)
    return candidate


@dataclass(frozen=True)
class Column:
    name: str
    source_name: str
    dtype: str

    def to_dict(self) -> dict:
        return {"name": self.name, "source_name": self.source_name, "dtype": self.dtype}


@dataclass(frozen=True)
class Table:
    name: str
    source_sheet: str
    row_count: int
    columns: tuple[Column, ...]
    parquet_path: str

    @property
    def column_names(self) -> list[str]:
        return [column.name for column in self.columns]

    def column(self, name: str) -> Column | None:
        return next((column for column in self.columns if column.name == name), None)

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "source_sheet": self.source_sheet,
            "row_count": self.row_count,
            "parquet": Path(self.parquet_path).name,
            "columns": [column.to_dict() for column in self.columns],
        }


@dataclass
class Workbook:
    """A workbook that has been converted to Parquet and is ready for DuckDB."""

    source_path: str
    file_name: str
    digest: str
    tables: tuple[Table, ...]
    cache: WorkbookCache
    from_cache: bool = False
    ingest_ms: int = 0
    _connections: list[duckdb.DuckDBPyConnection] = field(default_factory=list, repr=False)

    def table(self, name: str) -> Table | None:
        return next((table for table in self.tables if table.name == name), None)

    @property
    def largest_table(self) -> Table:
        return max(self.tables, key=lambda table: table.row_count)

    @property
    def total_rows(self) -> int:
        return sum(table.row_count for table in self.tables)

    def connect(self, *, memory_limit: str | None = None, threads: int | None = None,
                temp_directory: str | None = None) -> duckdb.DuckDBPyConnection:
        """Open an in-process DuckDB connection with one view per sheet."""
        connection = duckdb.connect(database=":memory:")
        if memory_limit:
            connection.execute(f"SET memory_limit={quote_literal(memory_limit)}")
        if threads:
            connection.execute(f"SET threads={int(threads)}")
        connection.execute(f"SET temp_directory={quote_literal(temp_directory or str(self.cache.directory))}")
        for table in self.tables:
            connection.execute(
                f"CREATE OR REPLACE VIEW {quote_identifier(table.name)} AS "
                f"SELECT * FROM read_parquet({quote_literal(table.parquet_path)})"
            )
        self._connections.append(connection)
        return connection

    def close(self) -> None:
        for connection in self._connections:
            try:
                connection.close()
            except Exception:  # pragma: no cover - connection already gone
                logger.debug("connection already closed", exc_info=True)
        self._connections.clear()

    def to_dict(self) -> dict:
        return {"source_name": self.file_name, "tables": [table.to_dict() for table in self.tables]}


# --------------------------------------------------------------------------- reading
def _normalise(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return str(value)


def _excel_sheets(path: Path, chunk_rows: int, max_rows: int) -> Iterator[tuple[str, list[str], Iterator[list[list[str | None]]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            while header is not None and all(cell is None for cell in header):
                header = next(rows, None)
            if header is None:
                logger.info("skipping empty sheet %s", worksheet.title)
                continue
            headers = [_normalise(cell) or f"column_{index + 1}" for index, cell in enumerate(header)]
            yield worksheet.title, headers, _chunk_rows(rows, len(headers), chunk_rows, max_rows, worksheet.title)
    finally:
        workbook.close()


def _chunk_rows(rows: Iterable[Sequence[object]], width: int, chunk_rows: int, max_rows: int,
                sheet: str) -> Iterator[list[list[str | None]]]:
    chunk: list[list[str | None]] = []
    seen = 0
    for row in rows:
        values = [_normalise(cell) for cell in row[:width]]
        if len(values) < width:
            values.extend([None] * (width - len(values)))
        if all(value is None for value in values):
            continue
        seen += 1
        if seen > max_rows:
            raise IngestError(f"Sheet '{sheet}' exceeds the {max_rows} row ingest limit")
        chunk.append(values)
        if len(chunk) >= chunk_rows:
            yield chunk
            chunk = []
    if chunk:
        yield chunk


def _write_text_parquet(target: Path, columns: list[str], chunks: Iterator[list[list[str | None]]]) -> int:
    """Stream chunks of text cells into a Parquet file; returns the row count."""
    schema = pa.schema([(ROW_ID, pa.int64())] + [(name, pa.string()) for name in columns])
    written = 0
    writer = pq.ParquetWriter(str(target), schema, compression="zstd")
    try:
        for chunk in chunks:
            arrays = [pa.array(range(written, written + len(chunk)), type=pa.int64())]
            arrays.extend(
                pa.array([row[index] for row in chunk], type=pa.string())
                for index in range(len(columns))
            )
            writer.write_table(pa.Table.from_arrays(arrays, schema=schema))
            written += len(chunk)
    finally:
        writer.close()
    return written


# --------------------------------------------------------------------------- typing
_TYPE_CANDIDATES = ("BIGINT", "DOUBLE", "BOOLEAN", "TIMESTAMP")


def _column_statistics(connection: duckdb.DuckDBPyConnection, source_sql: str,
                       columns: list[str]) -> tuple[int, dict[str, dict[str, int]]]:
    """One pass over the text data collecting cast statistics for every column.

    Casts alone are not enough: DuckDB happily truncates ``'12.7'`` to a BIGINT,
    and identifiers such as ``'007'`` would silently lose their leading zeros, so
    the shape of the text is checked as well.
    """
    selects = ["count(*) AS total_rows"]
    for index, column in enumerate(columns):
        quoted = quote_identifier(column)
        selects.extend([
            f"count({quoted}) AS non_null_{index}",
            f"count(*) FILTER (WHERE regexp_matches({quoted}, '^[+-]?[0-9]{{1,18}}$') "
            f"AND TRY_CAST({quoted} AS BIGINT) IS NOT NULL) AS bigint_{index}",
            f"count(TRY_CAST({quoted} AS DOUBLE)) AS double_{index}",
            f"count(*) FILTER (WHERE lower({quoted}) IN ('true','false')) AS boolean_{index}",
            f"count(TRY_CAST({quoted} AS TIMESTAMP)) AS timestamp_{index}",
            f"count(*) FILTER (WHERE TRY_CAST({quoted} AS TIMESTAMP) IS NOT NULL "
            f"AND TRY_CAST({quoted} AS TIMESTAMP) <> date_trunc('day', TRY_CAST({quoted} AS TIMESTAMP))) AS clocked_{index}",
            f"count(*) FILTER (WHERE regexp_matches({quoted}, '^[+-]?0[0-9]') "
            f"OR (length({quoted}) > 18 AND regexp_matches({quoted}, '^[+-]?[0-9]+$'))) AS opaque_{index}",
        ])
    row = connection.execute(f"SELECT {', '.join(selects)} FROM ({source_sql})").fetchone()
    assert row is not None
    total = int(row[0])
    statistics: dict[str, dict[str, int]] = {}
    for index, column in enumerate(columns):
        offset = 1 + index * 7
        statistics[column] = {
            "non_null": int(row[offset]), "bigint": int(row[offset + 1]), "double": int(row[offset + 2]),
            "boolean": int(row[offset + 3]), "timestamp": int(row[offset + 4]),
            "clocked": int(row[offset + 5]), "opaque": int(row[offset + 6]),
        }
    return total, statistics


def _decide_type(stats: dict[str, int]) -> str:
    """Pick the narrowest type that loses no value; text wins ties."""
    non_null = stats["non_null"]
    if non_null == 0:
        return "VARCHAR"
    if stats["opaque"]:
        return "VARCHAR"
    if stats["bigint"] == non_null:
        return "BIGINT"
    if stats["double"] == non_null:
        return "DOUBLE"
    if stats["boolean"] == non_null:
        return "BOOLEAN"
    if stats["timestamp"] == non_null:
        return "TIMESTAMP" if stats["clocked"] else "DATE"
    return "VARCHAR"


def _write_typed_parquet(connection: duckdb.DuckDBPyConnection, source_sql: str, columns: list[str],
                         target: Path) -> tuple[int, dict[str, str]]:
    total, statistics = _column_statistics(connection, source_sql, columns)
    types = {column: _decide_type(statistics[column]) for column in columns}
    projections = [quote_identifier(ROW_ID)]
    for column in columns:
        quoted = quote_identifier(column)
        dtype = types[column]
        projections.append(quoted if dtype == "VARCHAR" else f"TRY_CAST({quoted} AS {dtype}) AS {quoted}")
    connection.execute(
        f"COPY (SELECT {', '.join(projections)} FROM ({source_sql})) "
        f"TO {quote_literal(str(target))} (FORMAT PARQUET, COMPRESSION ZSTD)"
    )
    return total, types


# --------------------------------------------------------------------------- ingest
def _csv_source_sql(path: Path) -> tuple[str, list[str], list[str]]:
    """Return (sql, sanitized columns, source columns) for a delimited text file."""
    reader = (
        f"read_csv({quote_literal(str(path))}, all_varchar=true, header=true, "
        f"sample_size=-1, ignore_errors=false)"
    )
    with duckdb.connect(database=":memory:") as probe:
        described = probe.execute(f"DESCRIBE SELECT * FROM {reader}").fetchall()
    source_columns = [str(row[0]) for row in described]
    if not source_columns:
        raise IngestError("The file does not contain any columns")
    taken: set[str] = set()
    columns = [sanitize_identifier(name, taken) for name in source_columns]
    projections = [f"(row_number() OVER ()) - 1 AS {quote_identifier(ROW_ID)}"]
    projections.extend(
        f"{quote_identifier(source)} AS {quote_identifier(target)}"
        for source, target in zip(source_columns, columns)
    )
    return f"SELECT {', '.join(projections)} FROM {reader}", columns, source_columns


def _parquet_columns(connection: duckdb.DuckDBPyConnection, path: Path) -> dict[str, str]:
    described = connection.execute(
        f"DESCRIBE SELECT * FROM read_parquet({quote_literal(str(path))})"
    ).fetchall()
    return {str(row[0]): str(row[1]) for row in described}


def _load_cached(cache: WorkbookCache, source: Path) -> Workbook | None:
    manifest = cache.read_manifest()
    if not manifest:
        return None
    tables = []
    for entry in manifest.get("tables", []):
        parquet = cache.directory / str(entry.get("parquet", ""))
        if not parquet.exists():
            return None
        tables.append(Table(
            name=str(entry["name"]), source_sheet=str(entry.get("source_sheet", entry["name"])),
            row_count=int(entry.get("row_count", 0)),
            columns=tuple(Column(str(column["name"]), str(column.get("source_name", column["name"])),
                                 str(column.get("dtype", "VARCHAR"))) for column in entry.get("columns", [])),
            parquet_path=str(parquet),
        ))
    if not tables:
        return None
    return Workbook(source_path=str(source), file_name=str(manifest.get("source_name", source.name)),
                    digest=cache.digest, tables=tuple(tables), cache=cache, from_cache=True)


def ingest(path: str | os.PathLike[str], *, cache_root: str | os.PathLike[str] = "./data/cache",
           max_rows: int = DEFAULT_MAX_ROWS, chunk_rows: int = DEFAULT_CHUNK_ROWS,
           display_name: str | None = None, refresh: bool = False) -> Workbook:
    """Convert a workbook to cached Parquet tables and describe them.

    ``display_name`` keeps the original upload name in table names and in the
    schema card even when the file is stored on disk under a generated name.
    """
    source = Path(path)
    if not source.is_file():
        raise IngestError(f"File not found: {source}")
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise IngestError(f"Unsupported file type '{suffix}'. Use .xlsx, .xlsm or .csv")

    digest = file_digest(source)
    cache = WorkbookCache(root=Path(cache_root), digest=digest)
    if refresh:
        cache.clear()
    else:
        cached = _load_cached(cache, source)
        if cached:
            logger.info("workbook cache hit digest=%s tables=%d", digest[:12], len(cached.tables))
            return cached

    started = time.perf_counter()
    cache.prepare()
    tables: list[Table] = []
    taken_tables: set[str] = set()
    with duckdb.connect(database=":memory:") as connection:
        connection.execute(f"SET temp_directory={quote_literal(str(cache.directory))}")
        if suffix in CSV_SUFFIXES:
            source_sql, columns, source_columns = _csv_source_sql(source)
            stem = Path(display_name).stem if display_name else source.stem
            name = sanitize_identifier(stem, taken_tables, fallback="sheet")
            tables.append(_materialise(connection, cache, name, stem, source_sql, columns, source_columns))
        else:
            for sheet, headers, chunks in _excel_sheets(source, chunk_rows, max_rows):
                taken_columns: set[str] = set()
                columns = [sanitize_identifier(header, taken_columns) for header in headers]
                name = sanitize_identifier(sheet, taken_tables, fallback="sheet")
                raw_path = cache.directory / f"{safe_name(name)}.raw.parquet"
                try:
                    rows = _write_text_parquet(raw_path, columns, chunks)
                    if rows == 0:
                        logger.info("skipping sheet %s: no data rows", sheet)
                        continue
                    raw_sql = f"SELECT * FROM read_parquet({quote_literal(str(raw_path))})"
                    tables.append(_materialise(connection, cache, name, sheet, raw_sql, columns, headers))
                finally:
                    raw_path.unlink(missing_ok=True)
        if not tables:
            raise IngestError("The workbook does not contain any readable rows")

    workbook = Workbook(source_path=str(source), file_name=display_name or source.name, digest=digest,
                        tables=tuple(tables), cache=cache,
                        ingest_ms=int((time.perf_counter() - started) * 1000))
    cache.write_manifest({**workbook.to_dict(), "ingest_ms": workbook.ingest_ms})
    logger.info("ingested %s digest=%s rows=%d in %d ms", source.name, digest[:12],
                workbook.total_rows, workbook.ingest_ms)
    return workbook


def _materialise(connection: duckdb.DuckDBPyConnection, cache: WorkbookCache, name: str, sheet: str,
                 source_sql: str, columns: list[str], source_columns: Sequence[str]) -> Table:
    target = cache.parquet_path(name)
    row_count, _ = _write_typed_parquet(connection, source_sql, columns, target)
    dtypes = _parquet_columns(connection, target)
    return Table(
        name=name, source_sheet=sheet, row_count=row_count,
        columns=tuple(Column(column, str(source), dtypes.get(column, "VARCHAR"))
                      for column, source in zip(columns, source_columns)),
        parquet_path=str(target),
    )
