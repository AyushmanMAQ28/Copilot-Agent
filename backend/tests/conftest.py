"""Shared fixtures, including the synthetic 100k row workbook used by the QA tests."""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

REGIONS = ("North", "South", "East", "West")
CHANNELS = ("Web", "Store", "Partner")
LARGE_ROWS = 100_000

NOTE_SUBJECTS = ("delivery", "packaging", "invoice", "installation", "warranty")
NOTE_MOODS = ("praised", "questioned", "escalated", "confirmed", "disputed")


def _note(index: int) -> str:
    subject = NOTE_SUBJECTS[index % len(NOTE_SUBJECTS)]
    mood = NOTE_MOODS[(index // 7) % len(NOTE_MOODS)]
    return (f"Customer {mood} the {subject} experience for order {index} and asked the "
            f"support team to review the {subject} notes recorded in ticket {index % 977}.")


def large_workbook_facts() -> dict:
    """Ground truth for the synthetic workbook, computed in plain Python."""
    totals = {region: 0.0 for region in REGIONS}
    counts = {channel: 0 for channel in CHANNELS}
    amount_total = 0.0
    units_total = 0
    for index in range(LARGE_ROWS):
        amount = round(10 + (index % 977) * 1.25, 2)
        totals[REGIONS[index % len(REGIONS)]] += amount
        counts[CHANNELS[index % len(CHANNELS)]] += 1
        amount_total += amount
        units_total += index % 13
    return {
        "row_count": LARGE_ROWS,
        "amount_by_region": {region: round(value, 2) for region, value in totals.items()},
        "rows_by_channel": counts,
        "amount_total": round(amount_total, 2),
        "units_total": units_total,
    }


def write_large_workbook(path: Path) -> Path:
    """A deterministic 1 lakh row workbook plus a small lookup sheet."""
    from datetime import date, timedelta

    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    sales = workbook.create_sheet("Sales")
    sales.append(["order_id", "order_date", "region", "channel", "units", "amount", "note"])
    start = date(2024, 1, 1)
    for index in range(LARGE_ROWS):
        sales.append([
            f"ORD-{index:06d}",
            start + timedelta(days=index % 366),
            REGIONS[index % len(REGIONS)],
            CHANNELS[index % len(CHANNELS)],
            index % 13,
            round(10 + (index % 977) * 1.25, 2),
            _note(index),
        ])
    regions = workbook.create_sheet("Regions")
    regions.append(["region", "manager"])
    for position, region in enumerate(REGIONS):
        regions.append([region, f"Manager {position}"])
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture(scope="session")
def large_workbook(tmp_path_factory) -> Path:
    directory = tmp_path_factory.mktemp("large-workbook")
    return write_large_workbook(directory / "sales_100k.xlsx")


@pytest.fixture(scope="session")
def large_cache(tmp_path_factory) -> Path:
    return tmp_path_factory.mktemp("qa-cache")


@pytest.fixture(scope="session")
def large_engine(large_workbook: Path, large_cache: Path):
    from app.qa import QAEngine

    engine = QAEngine(large_workbook, cache_root=large_cache)
    yield engine
    engine.close()


@pytest.fixture
def small_workbook(tmp_path: Path) -> Path:
    """A tiny two sheet workbook with mixed column types."""
    from datetime import date

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Q1 Sales"
    sheet.append(["Order ID", "Order Date", "Region", "Units", "Amount", "Comment"])
    rows = [
        ("A-1", date(2024, 1, 5), "North", 3, 120.55, "Repeat buyer asked about warranty coverage details."),
        ("A-2", date(2024, 1, 9), "South", 1, 80.10, "New buyer needed installation guidance for the kit."),
        ("A-3", date(2024, 2, 2), "North", 7, 940.00, "Bulk order shipped early with a discounted freight rate."),
        ("A-4", date(2024, 2, 20), "West", 2, 45.25, "Buyer disputed the invoice total before paying online."),
    ]
    for row in rows:
        sheet.append(row)
    lookup = workbook.create_sheet("Regions")
    lookup.append(["region", "manager"])
    for region, manager in (("North", "Ada"), ("South", "Grace"), ("West", "Linus")):
        lookup.append([region, manager])
    target = tmp_path / "small.xlsx"
    workbook.save(target)
    workbook.close()
    return target
